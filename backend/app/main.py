
from starlette.concurrency import run_in_threadpool
from app.services.import_service import save_analysis
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader
from app.services.vocabulary_service import analyze_vocabulary

from app.dictionary_routes import router as dictionary_router

from app.library_routes import router as library_router
from app.enrichment_routes import router as enrichment_router

from app.study_routes import router as study_router

app = FastAPI(title="Learn English API")
app.include_router(dictionary_router)
app.include_router(library_router)
app.include_router(enrichment_router)
app.include_router(study_router)

MAX_FILE_SIZE = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc"}

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "https://learn-english-khaki.vercel.app"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Retry-After"],
)


@app.get("/health")
def health():
    return {"status": "ok"}


def extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    return "\n".join(
        page.extract_text() or ""
        for page in reader.pages
    )


def extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    lines = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            lines.append(
                " | ".join(cell.text for cell in row.cells)
            )

    return "\n".join(lines)


@app.post("/analyze")
async def analyze(
    text: str | None = Form(default=None),
    file: UploadFile | None = File(default=None),
):
    if file is not None and text and text.strip():
        raise HTTPException(
            status_code=400,
            detail="Chỉ nhập văn bản hoặc tải file.",
        )

    if file is None and not (text and text.strip()):
        raise HTTPException(
            status_code=400,
            detail="Vui lòng nhập nội dung hoặc tải file.",
        )

    if file is not None:
        filename = file.filename or ""
        extension = Path(filename).suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail="Định dạng file không được hỗ trợ.",
            )

        if extension == ".doc":
            raise HTTPException(
                status_code=400,
                detail="File DOC cũ chưa được hỗ trợ. "
                       "Vui lòng chuyển thành DOCX.",
            )

        # Đọc tối đa 10 MB + 1 byte để phát hiện file quá lớn
        content = await file.read(MAX_FILE_SIZE + 1)

        if len(content) > MAX_FILE_SIZE:
            raise HTTPException(
                status_code=413,
                detail="File vượt quá 10 MB.",
            )

        try:
            if extension == ".pdf":
                extracted_text = extract_pdf(content)
            else:
                extracted_text = extract_docx(content)
        except Exception as exc:
            raise HTTPException(
                status_code=422,
                detail="Không thể đọc nội dung file.",
            ) from exc

        source = "file"

    else:
        extracted_text = text or ""
        filename = None
        source = "text"

    extracted_text = extracted_text.strip()

    if not extracted_text:
        raise HTTPException(
            status_code=422,
            detail="Không tìm thấy văn bản trong tài liệu.",
        )
    try:
        vocabulary_result = await run_in_threadpool(analyze_vocabulary, extracted_text)
    except ValueError as exc:
        raise HTTPException(
            status_code=413,
            detail=str(exc),
        ) from exc
    storage = await run_in_threadpool(save_analysis, vocabulary_result)
    return {
        "storage": storage,
        "source": source,
        "filename": filename,
        "text": extracted_text,
        "character_count": len(extracted_text),
        "analysis": vocabulary_result,
    }
